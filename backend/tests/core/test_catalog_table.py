"""The catalog as a plain table (`app.cli catalog-table`): a view of the catalog, one row each."""

from pathlib import Path

from openpyxl import load_workbook

from app.core.config import get_settings
from app.services.engine.reference import Size
from app.services.quick.config import quick_config
from app.services.seeding.catalog_table import HEADERS, rows, write_excel, write_markdown
from app.services.seeding.reference import load_reference


def test_one_row_per_block_with_its_days_and_uses(tmp_path: Path) -> None:
    catalog = load_reference(get_settings().seeds_dir).catalog
    table = rows(catalog, quick_config())
    assert sorted(r.key for r in table) == sorted(c.key for c in catalog.components)
    user = next(r for r in table if r.key == "user_management_rbac")
    block = next(c for c in catalog.components if c.key == "user_management_rbac")
    assert (user.small, user.large) == (block.base_effort_pd[Size.S], block.base_effort_pd[Size.L])
    assert "Standard block" in user.used_by
    write_markdown(catalog, table, tmp_path / "c.md")
    write_excel(catalog, table, tmp_path / "c.xlsx")
    assert "| User management and role-based access |" in (tmp_path / "c.md").read_text("utf-8")
    sheet = load_workbook(tmp_path / "c.xlsx").active
    assert sheet is not None and sheet.max_row == len(table) + 1
    assert [c.value for c in sheet[1]] == HEADERS
