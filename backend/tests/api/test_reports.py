"""Exports (brief M2): PDF report and Excel BoQ download cleanly and carry exactly the
numbers the API returns."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

H = dict[str, str]


def paise(cell: object) -> int:
    assert isinstance(cell, int | float), cell
    return round(cell * 100)


def _estimate(client: TestClient, headers: H) -> dict[str, object]:
    body = {"name": "Report portal", "features": ["online_forms", "payments", "ai_chatbot"]}
    s = client.post("/quick-estimates", json=body, headers=headers).json()
    detail: dict[str, object] = client.get(f"/estimates/{s['estimate_id']}", headers=headers).json()
    return detail


def test_pdf_report(client: TestClient, admin_headers: H) -> None:
    est = _estimate(client, admin_headers)
    r = client.get(f"/estimates/{est['id']}/report.pdf", headers=admin_headers)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF") and len(r.content) > 5000
    assert 'filename="Report-portal-v1.pdf"' in r.headers["content-disposition"]


def test_excel_boq_matches_api_exactly(client: TestClient, admin_headers: H) -> None:
    est = _estimate(client, admin_headers)
    r = client.get(f"/estimates/{est['id']}/boq.xlsx", headers=admin_headers)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Summary", "Cost lines", "By stage", "By year", "Manpower", "Assumptions"} <= set(
        wb.sheetnames
    )
    summary = {
        row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row and row[0]
    }
    assert (
        paise(summary["Likely price (P50: planned cost + risk allowance, with GST)"])
        == est["p50_paise"]
    )
    assert paise(summary["Low (P10)"]) == est["p10_paise"]
    rows = list(wb["Cost lines"].iter_rows(min_row=2, values_only=True))
    api_lines = est["lines"]
    assert isinstance(api_lines, list) and len(rows) == len(api_lines)
    for row, line in zip(rows, api_lines, strict=True):
        assert paise(row[6]) == line["likely_paise"], line["line_key"]
        assert paise(row[5]) == line["low_paise"] and paise(row[7]) == line["high_paise"]


def test_exports_are_org_scoped(client: TestClient, admin_headers: H) -> None:
    assert client.get("/estimates/not-an-id/report.pdf", headers=admin_headers).status_code == 404
    assert client.get("/estimates/not-an-id/boq.xlsx").status_code == 401
