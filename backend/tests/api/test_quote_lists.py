"""What a quote rests on and leaves out (OPEN_QUESTIONS 147): answers nobody gave are listed as
assumed, things marked "Not included" or unticked are not priced and are listed, and a "Not
sure" about old data or other software leaves it out with an "if required" price. The PDF and
Excel print the same lists."""

import io
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

H = dict[str, str]
BASE = {"name": "Quote lists", "features": ["online_forms", "ai_chatbot"]}


def quick(client: TestClient, headers: H, **body: Any) -> Any:
    return client.post("/quick-estimates", json={**BASE, **body}, headers=headers)


def test_untouched_answers_are_listed_as_assumed(client: TestClient, admin_headers: H) -> None:
    s = quick(client, admin_headers, answered=[]).json()
    assumed = {a["key"]: a for a in s["assumed"]}
    assert assumed["scale"]["detail"] == "10,000 users, 200 at the same time"
    assert "availability" in assumed  # never asked on the form: still a silent default
    assert "ai_passthrough" not in assumed  # decided on every estimate
    assert "onsite_period" not in assumed  # no staff at the client's office
    touched = quick(client, admin_headers, answered=["scale"]).json()
    assert "scale" not in {a["key"] for a in touched["assumed"]}
    # older forms send no list: an answer other than the default was given
    old = quick(client, admin_headers, choices={"scale": "up_to_1k"}).json()
    assert "scale" not in {a["key"] for a in old["assumed"]}
    assert "version" in {a["key"] for a in old["assumed"]}


def test_client_type_is_not_recorded(client: TestClient, admin_headers: H) -> None:
    s = quick(client, admin_headers).json()
    detail = client.get(f"/estimates/{s['estimate_id']}", headers=admin_headers).json()
    assert "buyer_segment" not in {a["key"] for a in detail["assumptions"]}


def test_feature_marked_not_included(client: TestClient, admin_headers: H) -> None:
    plain = quick(client, admin_headers).json()
    out = {"kind": "feature", "key": "payments", "quote": "no online payments"}
    s = quick(client, admin_headers, not_included=[out]).json()
    assert s["build_only_paise"] == plain["build_only_paise"]  # it was never priced
    item = next(i for i in s["not_included"] if i["why"] == "excluded_by_you")
    assert item["quote"] == "no online payments" and "payment" in item["label"].lower()
    both = quick(client, admin_headers, features=["online_forms", "payments"], not_included=[out])
    assert both.status_code == 422  # priced and not included at once


def test_standard_block_can_be_unticked(client: TestClient, admin_headers: H) -> None:
    blocks = client.post("/quick-estimate/standard-blocks", json=BASE, headers=admin_headers)
    assert blocks.status_code == 200, blocks.text
    keys = {b["key"]: b for b in blocks.json()}
    assert keys["devops_cicd_setup"]["label"] == "Servers and release set-up"
    plain = quick(client, admin_headers).json()
    out = {"kind": "standard", "key": "devops_cicd_setup"}
    s = quick(client, admin_headers, not_included=[out]).json()
    assert s["build_only_paise"] < plain["build_only_paise"]
    removed = [i for i in s["not_included"] if i["why"] == "removed_by_you"]
    assert [i["label"] for i in removed] == ["Servers and release set-up"]
    bad = quick(client, admin_headers, not_included=[{"kind": "standard", "key": "payments"}])
    assert bad.status_code == 422


def test_not_sure_about_old_data_is_left_out_and_quoted(
    client: TestClient, admin_headers: H
) -> None:
    plain = quick(client, admin_headers).json()
    s = quick(client, admin_headers, choices={"old_data": "not_sure"}).json()
    assert s["build_only_paise"] == plain["build_only_paise"]  # not in the price
    assert "Move data from the old system" not in s["features"]
    assert s["not_sure"] == []  # left out, so it does not widen the range
    item = next(i for i in s["not_included"] if i["why"] == "not_sure")
    assert item["question_key"] == "old_data"
    options = client.get(f"/estimates/{s['estimate_id']}/options", headers=admin_headers).json()
    old = next(o for o in options if o["key"] == "old_data")
    assert old["why"] == "left_out"
    # if required: a little to a lot; the usual size (as "Not sure" priced it before) between
    assert 0 < old["low_change_paise"] < old["usual_change_paise"] < old["high_change_paise"]


def test_description_rules_features_out(client: TestClient, admin_headers: H) -> None:
    text = "A booking app for a salon, no mobile app needed. Payments not needed."
    r = client.post("/quick-estimate/suggest", json={"text": text}, headers=admin_headers)
    assert r.status_code == 200, r.text
    body = r.json()
    excluded = {f["key"]: f["quote"] for f in body["excluded"]}
    assert excluded.keys() >= {"mobile_app", "payments"}
    assert "no mobile app" in excluded["mobile_app"]
    assert "payments" not in {f["key"] for f in body["features"]}


def test_reports_print_assumed_and_not_included(client: TestClient, admin_headers: H) -> None:
    s = quick(
        client,
        admin_headers,
        answered=[],
        choices={"old_data": "not_sure"},
        not_included=[{"kind": "standard", "key": "devops_cicd_setup"}],
    ).json()
    r = client.get(f"/estimates/{s['estimate_id']}/boq.xlsx", headers=admin_headers)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Assumed answers", "Not included"} <= set(wb.sheetnames)
    rows = list(wb["Not included"].iter_rows(min_row=2, values_only=True))
    why = {str(row[0]): str(row[1]) for row in rows}
    assert "If required:" in why["Move data from the old system"]
    assert "Servers and release set-up" in why
    assumed = [row[0] for row in wb["Assumed answers"].iter_rows(min_row=2, values_only=True)]
    assert "How many people will use it?" in assumed
    keys = [row[0] for row in wb["Assumptions"].iter_rows(min_row=2, values_only=True)]
    assert "buyer_segment" not in keys
    pdf = client.get(f"/estimates/{s['estimate_id']}/report.pdf", headers=admin_headers)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
