# ruff: noqa: E501  (inline fixture data)
"""Tender upload: parse (PDF/DOCX), read (basic reader without a key, or AI via a fake client),
verify every citation, then price the confirmed form with page-level traceability."""

import io
from collections.abc import Iterator
from typing import Any

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.llm.client import FakeClient, set_llm_client
from app.models import Organisation, Role, User
from tests.conftest import TEST_PASSWORD, login

H = dict[str, str]

TENDER_PAGES = [
    [
        "Request for Proposal for Online Vendor Empanelment Portal",
        "Issued by Demo Public Sector Undertaking",
        "2.4 The Earnest Money Deposit (EMD) of Rs. 5,00,000 shall be submitted with the bid.",
        "2.6 The bidder shall furnish a Performance Bank Guarantee of 3% of the contract value.",
        "2.9 Go-live shall be achieved within 6 months of the issue of work order.",
    ],
    [
        "4.1 Vendors shall submit an online application with supporting documents.",
        "4.2 Applications shall pass a multi-level approval workflow with SLA tracking.",
        "4.3 Fees shall be paid through the payment gateway of the PSU.",
        "4.4 An AI chatbot shall answer vendor queries in English and Hindi.",
        "5.1 The portal will serve 50,000 registered users with 500 concurrent users at peak.",
        "5.2 The system shall be hosted on a MeitY empanelled cloud with 99.9% availability.",
        "8.1 The bidder shall provide 3 years of operations and maintenance after go-live.",
        "8.4 Training shall be provided to officials as per requirement.",
    ],
]


def make_docx(pages: list[list[str]]) -> bytes:
    document = docx.Document()
    for page in pages:
        for line in page:
            document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal, valid text PDF (Helvetica, one line per row) — no PDF library needed."""
    objects: list[bytes] = []
    n_pages = len(pages)
    page_ids = [3 + 2 * i for i in range(n_pages)]
    font_id = 3 + 2 * n_pages
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids).encode()
    objects.append(
        b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(n_pages).encode() + b" >>"
    )
    for i, lines in enumerate(pages):
        content = (
            "BT /F1 9 Tf 40 800 Td 12 TL "
            + " ".join(
                "(" + ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") '"
                for ln in lines
            )
            + " ET"
        )
        stream = content.encode("latin-1")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {page_ids[i] + 1} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for num, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{num} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()


def upload(client: TestClient, headers: H, name: str, data: bytes) -> Any:
    return client.post("/documents", files={"file": (name, data)}, headers=headers)


@pytest.fixture
def no_override() -> Iterator[None]:
    set_llm_client(None)
    yield
    set_llm_client(None)


def test_basic_reader_without_a_key(
    client: TestClient, admin_headers: H, no_override: None
) -> None:
    r = upload(client, admin_headers, "tender.pdf", make_pdf(TENDER_PAGES))
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["status"] == "extracted" and doc["pages"] == 2 and not doc["approximate_pages"]
    ex = doc["extraction"]
    assert ex["reader"] == "basic"
    keys = {f["key"] for f in ex["features"]}
    assert {"online_forms", "approval_workflow", "payments", "ai_chatbot", "languages"} <= keys
    pay = next(f for f in ex["features"] if f["key"] == "payments")
    assert pay["evidence"][0] == {
        "page": 2,
        "clause_ref": "4.3",
        "quote": "4.3 Fees shall be paid through the payment gateway of the PSU.",
        "document": "tender.pdf",
        "checked": True,
    }
    n = ex["numbers"]
    assert n["om_years"]["value"] == 3 and n["deadline_months"]["value"] == 6
    assert n["users_total"]["value"] == 50000 and n["peak_concurrent_users"]["value"] == 500
    assert n["emd_paise"]["value"] == 50_000_000 and n["pbg_percent"]["value"] == 0.03
    c = ex["choices"]
    assert c["availability"]["value"] == "high" and c["hosting"]["value"] == "indian_cloud"
    assert c["scale"]["value"] == "up_to_1l"  # 50,000 users
    assert doc["project_name"].startswith("Request for Proposal")


def test_docx_upload_has_approximate_pages(
    client: TestClient, admin_headers: H, no_override: None
) -> None:
    doc = upload(client, admin_headers, "tender.docx", make_docx(TENDER_PAGES)).json()
    assert doc["file_type"] == "docx" and doc["approximate_pages"]
    assert {f["key"] for f in doc["extraction"]["features"]} >= {"payments", "ai_chatbot"}


def test_ai_reader_citations_are_verified(client: TestClient, admin_headers: H) -> None:
    """The AI path: invented quotes and unknown keys are dropped, real ones are kept."""

    def responder(system: str, user: str) -> object:
        assert "<tender_document>" in user and 'page number="2"' in user
        assert "Never estimate" in system  # rule 1 is in the prompt
        return {
            "project_name": "Online Vendor Empanelment Portal",
            "summary": "A portal for vendors to apply online.",
            "features": [
                {"key": "payments", "evidence": [{"page": 2, "clause_ref": "4.3", "quote": "Fees shall be paid through the payment gateway"}]},
                {"key": "aadhaar_digilocker", "evidence": [{"page": 2, "clause_ref": "4.9", "quote": "Citizens shall log in with Aadhaar"}]},
                {"key": "teleport", "evidence": [{"page": 1, "clause_ref": "", "quote": "Request for Proposal for Online"}]},
            ],
            "other_features": [
                {"name": "Vehicle tracking", "description": "Track inspection vehicles on a map.", "category": "module", "size": "M", "evidence": [{"page": 2, "clause_ref": "4.2", "quote": "multi-level approval workflow with SLA tracking"}]},
                {"name": "Made up thing", "description": "x", "category": "module", "size": "M", "evidence": [{"page": 2, "clause_ref": "", "quote": "this sentence is not in the document"}]},
                {"name": "Bad type", "description": "x", "category": "spaceship", "size": "M", "evidence": [{"page": 2, "clause_ref": "4.1", "quote": "Vendors shall submit an online application"}]},
            ],
            "choices": [{"question": "support_hours", "value": "always", "evidence": [{"page": 2, "clause_ref": "5.2", "quote": "99.9% availability"}]}],
            "numbers": [{"field": "om_years", "value": 3, "evidence": [{"page": 2, "clause_ref": "8.1", "quote": "3 years of operations and maintenance"}]}],
            "unclear_points": [{"topic": "Helpdesk hours", "why": "Not stated", "page": None, "clause_ref": "", "quote": "", "prebid_query": "Kindly confirm the helpdesk support hours."}],
        }  # fmt: skip

    set_llm_client(FakeClient(responder))
    try:
        doc = upload(client, admin_headers, "tender.pdf", make_pdf(TENDER_PAGES)).json()
    finally:
        set_llm_client(None)
    ex = doc["extraction"]
    assert ex["reader"] == "ai" and doc["llm_model"] == "fake"
    assert [f["key"] for f in ex["features"]] == ["payments"]  # invented quote + bad key dropped
    assert ex["choices"]["support_hours"]["value"] == "always"  # quote is real, user will review
    assert ex["numbers"]["om_years"]["value"] == 3
    assert ex["dropped_findings"] == 4  # + invented and wrongly typed other features
    assert [o["name"] for o in ex["other_features"]] == ["Vehicle tracking"]
    assert ex["other_features"][0]["evidence"][0]["clause_ref"] == "4.2"
    assert ex["unclear_points"][0]["topic"] == "Helpdesk hours"


def test_estimate_from_document_traces_to_pages(
    client: TestClient, admin_headers: H, no_override: None
) -> None:
    doc = upload(client, admin_headers, "tender.pdf", make_pdf(TENDER_PAGES)).json()
    ex = doc["extraction"]
    body = {
        "name": "Vendor empanelment portal",
        "document_id": doc["id"],
        "features": [f["key"] for f in ex["features"]],
        "choices": {k: v["value"] for k, v in ex["choices"].items() if k != "support_years"},
        "overrides": {k: v["value"] for k, v in ex["numbers"].items()},
    }
    s = client.post("/quick-estimates", json=body, headers=admin_headers)
    assert s.status_code == 201, s.text
    summary = s.json()
    assert summary["project_id"] == doc["project_id"] and summary["support_years"] == 3
    detail = client.get(f"/estimates/{summary['estimate_id']}", headers=admin_headers).json()
    assert detail["timeline"]["deadline_months"] == "6"
    traces = [
        client.get(
            f"/estimates/{summary['estimate_id']}/lines/{ln['id']}/trace", headers=admin_headers
        ).json()["trace"]
        for ln in detail["lines"]
        if ln["bucket"] == "build.integrations"
    ]
    payment = next(
        t for t in traces if t["component"]["catalog_key"] == "integration_payment_gateway"
    )
    pages = {(r["clause_ref"], r["page"]) for r in payment["requirements"]}
    assert ("4.3", 2) in pages  # payment gateway -> RFP clause 4.3 on page 2
    emd = next(ln for ln in detail["lines"] if ln["bucket"] == "commercial.emd")
    emd_trace = client.get(
        f"/estimates/{summary['estimate_id']}/lines/{emd['id']}/trace", headers=admin_headers
    ).json()["trace"]
    assert any(r["clause_ref"] == "2.4" and r["page"] == 1 for r in emd_trace["requirements"])
    # recalculating from the same upload replaces the mapping instead of duplicating it
    again = client.post("/quick-estimates", json=body, headers=admin_headers).json()
    assert (
        again["project_id"] == doc["project_id"]
        and again["planned_paise"] == summary["planned_paise"]
    )


def test_document_and_manual_flows_price_alike(
    client: TestClient, admin_headers: H, no_override: None
) -> None:
    """The same confirmed form priced from an uploaded document and typed in by hand: same
    components, same reference versions, the same price on every line. Only the citations
    differ (the document flow traces each requirement to its page)."""
    doc = upload(client, admin_headers, "tender.pdf", make_pdf(TENDER_PAGES)).json()
    ex = doc["extraction"]
    form = {
        "features": [f["key"] for f in ex["features"]],
        "choices": {k: v["value"] for k, v in ex["choices"].items()},
        "overrides": {k: v["value"] for k, v in ex["numbers"].items()},
    }
    runs = []
    for extra in ({"document_id": doc["id"]}, {}):
        s = client.post(
            "/quick-estimates", json={"name": "Portal", **form, **extra}, headers=admin_headers
        )
        assert s.status_code == 201, s.text
        detail = client.get(f"/estimates/{s.json()['estimate_id']}", headers=admin_headers).json()
        runs.append((s.json(), detail))
    (from_doc, doc_detail), (manual, manual_detail) = runs
    assert from_doc["project_id"] == doc["project_id"] != manual["project_id"]
    for key in ("catalog_version", "rate_card_version", "parameter_set_version"):
        assert doc_detail[key] == manual_detail[key]
    # component line keys carry each project's own component ids: compare what each line is
    lines = [
        sorted((ln["bucket"], ln["label"], ln["year"], ln["likely_paise"]) for ln in d["lines"])
        for _, d in runs
    ]
    assert lines[0] == lines[1]
    assert from_doc["planned_paise"] == manual["planned_paise"]
    assert from_doc["build_only_paise"] == manual["build_only_paise"]
    pages = {
        r.get("page")
        for ln in doc_detail["lines"]
        if ln["bucket"] == "build.integrations"
        for r in client.get(
            f"/estimates/{from_doc['estimate_id']}/lines/{ln['id']}/trace", headers=admin_headers
        ).json()["trace"]["requirements"]
    }
    assert 2 in pages  # cited from the document's page 2


SCOPE_PAGES = [
    *TENDER_PAGES,
    [
        "6.1 All officials shall sign in through two-factor authentication.",
        "6.2 The system shall keep an audit trail of every approval decision.",
    ],
]


def test_document_scope_is_suggested_and_priced_like_the_manual_form(
    client: TestClient, admin_headers: H, no_override: None
) -> None:
    """The document's own words suggest scope items (with their page); kept ticked, they are
    priced exactly as the same items chosen by hand."""
    doc = upload(client, admin_headers, "tender.pdf", make_pdf(SCOPE_PAGES)).json()
    ex = doc["extraction"]
    found = {s["key"]: s for s in ex["scope"]}
    assert {"two_step_sign_in", "workflow_audit"} <= set(found)
    assert found["two_step_sign_in"]["evidence"][0]["page"] == 3
    form = {
        "features": [f["key"] for f in ex["features"]],
        "choices": {k: v["value"] for k, v in ex["choices"].items()},
        "overrides": {k: v["value"] for k, v in ex["numbers"].items()},
        "scope": sorted(found),
    }
    runs = []
    for extra in ({"document_id": doc["id"]}, {}):
        s = client.post(
            "/quick-estimates", json={"name": "Portal", **form, **extra}, headers=admin_headers
        )
        assert s.status_code == 201, s.text
        detail = client.get(f"/estimates/{s.json()['estimate_id']}", headers=admin_headers).json()
        runs.append(
            sorted(
                (ln["bucket"], ln["label"], ln["year"], ln["likely_paise"])
                for ln in detail["lines"]
            )
        )
    assert runs[0] == runs[1]
    # found in the document but not kept: priced exactly like the form without them
    priced = []
    for extra in ({"document_id": doc["id"]}, {}):
        s = client.post(
            "/quick-estimates",
            json={"name": "Portal", **form, "scope": [], **extra},
            headers=admin_headers,
        ).json()
        assert s["scope"] == [] and s["scope_items"] == []
        detail = client.get(f"/estimates/{s['estimate_id']}", headers=admin_headers).json()
        priced.append(
            sorted(
                (ln["bucket"], ln["label"], ln["year"], ln["likely_paise"])
                for ln in detail["lines"]
            )
        )
    assert priced[0] == priced[1] != runs[0]


def test_prebid_queries(client: TestClient, admin_headers: H, no_override: None) -> None:
    doc = upload(client, admin_headers, "tender.pdf", make_pdf(TENDER_PAGES)).json()
    r = client.get(f"/documents/{doc['id']}/prebid-queries", headers=admin_headers)
    assert r.status_code == 200
    queries = r.json()
    vague = [q for q in queries if q["kind"] == "unclear"]
    assert vague and vague[0]["clause_ref"] == "8.4" and vague[0]["page"] == 2
    assert "8.4" in vague[0]["query"] and "as per requirement" in vague[0]["quote"]
    missing = {q["topic"] for q in queries if q["kind"] == "missing"}
    assert "Support hours not stated" in missing  # tender never says office hours or 24x7
    assert "Om years not stated" not in missing  # stated in clause 8.1
    assert [q["number"] for q in queries] == list(range(1, len(queries) + 1))


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        ("notes.txt", b"hello", "PDF or a Word"),
        ("fake.pdf", b"not a pdf at all", "PDF or a Word"),
        ("blank.pdf", make_pdf([[" "]]), "scanned"),
    ],
)
def test_bad_files_get_clear_messages(
    client: TestClient, admin_headers: H, name: str, data: bytes, message: str
) -> None:
    r = upload(client, admin_headers, name, data)
    assert r.status_code == 400 and message in r.json()["detail"]


def test_upload_is_org_scoped(
    client: TestClient, admin_headers: H, db: Session, no_override: None
) -> None:
    doc = upload(client, admin_headers, "tender.pdf", make_pdf(TENDER_PAGES)).json()
    other = Organisation(name="Other SI (documents)")
    db.add(other)
    db.flush()
    db.add(
        User(
            org_id=other.id,
            email="docs-spy@other.local",
            full_name="Spy",
            role=Role.ADMIN,
            password_hash=hash_password(TEST_PASSWORD),
        )
    )
    db.commit()
    spy = login(client, "docs-spy@other.local")
    assert client.get(f"/documents/{doc['id']}", headers=spy).status_code == 404
    assert client.get(f"/documents/{doc['id']}", headers=admin_headers).status_code == 200


def test_corrigendum_overrides_rfp(client: TestClient, admin_headers: H, no_override: None) -> None:
    rfp = upload(client, admin_headers, "rfp.pdf", make_pdf(TENDER_PAGES)).json()
    corr_pages = [
        [
            "Corrigendum No. 1",
            "2.9 Go-live shall be achieved within 7 months of the issue of work order.",
        ]
    ]
    r = client.post(
        "/documents",
        files={"file": ("corrigendum-1.pdf", make_pdf(corr_pages))},
        data={"project_id": rfp["project_id"], "kind": "corrigendum"},
        headers=admin_headers,
    )
    assert r.status_code == 201, r.text
    assert r.json()["project_id"] == rfp["project_id"]
    assert r.json()["project_name"] == rfp["project_name"]  # first document names the project
    merged = client.get(f"/projects/{rfp['project_id']}/extraction", headers=admin_headers).json()
    assert [d["kind"] for d in merged["documents"]] == ["rfp", "corrigendum"]
    assert merged["extraction"]["numbers"]["deadline_months"]["value"] == 7
    assert merged["changes"] == [
        {
            "field": "deadline_months",
            "old_value": 6,
            "new_value": 7,
            "document": "corrigendum-1.pdf",
            "kind": "corrigendum",
        }
    ]
    ex = merged["extraction"]
    body = {
        "name": "Portal",
        "document_id": rfp["id"],
        "features": [f["key"] for f in ex["features"]],
        "overrides": {"deadline_months": 7},
    }
    s = client.post("/quick-estimates", json=body, headers=admin_headers).json()
    detail = client.get(f"/estimates/{s['estimate_id']}", headers=admin_headers).json()
    assert detail["timeline"]["deadline_months"] == "7"
    pm = next(ln for ln in detail["lines"] if ln["line_key"] == "programme.project_management/pm")
    trace = client.get(
        f"/estimates/{s['estimate_id']}/lines/{pm['id']}/trace", headers=admin_headers
    ).json()["trace"]
    deadline = next(d for d in trace["drivers"] if d["key"] == "deadline_months")
    assert deadline["value"] == "7" and deadline["requirement_ids"]
    cited = [r for r in trace["requirements"] if r["id"] in deadline["requirement_ids"]]
    assert cited and cited[0]["clause_ref"] == "2.9"


def test_unknown_document_kind_rejected(client: TestClient, admin_headers: H) -> None:
    r = client.post(
        "/documents",
        files={"file": ("x.pdf", make_pdf(TENDER_PAGES))},
        data={"kind": "love-letter"},
        headers=admin_headers,
    )
    assert r.status_code == 422


def test_scanned_pdf_read_by_ai_with_unchecked_quotes(client: TestClient, admin_headers: H) -> None:
    seen: dict[str, str] = {}

    def responder(system: str, user: str) -> object:
        seen["user"] = user
        return {
            "project_name": "Scanned tender",
            "summary": "A scanned RFP.",
            "features": [{"key": "payments", "evidence": [{"page": 1, "clause_ref": "3.2", "quote": "Payment through gateway of the department"}]}],
            "other_features": [],
            "choices": [],
            "numbers": [{"field": "om_years", "value": 2, "evidence": [{"page": 9, "clause_ref": "", "quote": "two years of O&M after go-live"}]}],
            "unclear_points": [],
        }  # fmt: skip

    set_llm_client(FakeClient(responder))
    try:
        r = upload(client, admin_headers, "scanned.pdf", make_pdf([[" "], [" "]]))
    finally:
        set_llm_client(None)
    assert r.status_code == 201, r.text
    assert seen["user"].startswith("[PDF ")  # sent as a PDF, not as extracted text
    ex = r.json()["extraction"]
    assert ex["features"][0]["evidence"][0]["checked"] is False
    assert ex["numbers"] == {}  # page 9 does not exist in a 2-page document -> dropped


def test_basic_reader_ignores_paperwork_deadlines() -> None:
    from app.services.documents.parsing import Page, ParsedDocument
    from app.services.extraction.basic_reader import basic_read
    from app.services.quick.config import quick_config

    def deadline(text: str) -> float | None:
        doc = ParsedDocument(kind="pdf", pages=[Page(number=1, text=text)], approximate_pages=False)
        found = [n.value for n in basic_read(quick_config(), doc).numbers]
        fields = [n.field for n in basic_read(quick_config(), doc).numbers]
        return found[fields.index("deadline_months")] if "deadline_months" in fields else None

    assert deadline("The bidder shall submit the PBG within 15 days of award of contract.") is None
    assert deadline("Development shall be completed within 45 days of award of work.") == 2


def test_long_documents_offer_but_do_not_tick_passing_mentions() -> None:
    from app.services.documents.parsing import Page, ParsedDocument
    from app.services.extraction.basic_reader import basic_read
    from app.services.quick.config import quick_config

    cfg = quick_config()
    filler = [Page(number=n, text="General terms and conditions.") for n in range(2, 30)]
    once = [Page(number=1, text="Bids are shown on a dashboard.")]
    often = [Page(number=n, text=f"{n}.1 The dashboard shows analytics.") for n in (1, 2, 3)]

    def feature(pages: list[Page]) -> tuple[bool, int]:
        doc = ParsedDocument(kind="pdf", pages=pages, approximate_pages=False)
        f = next(x for x in basic_read(cfg, doc).features if x.key == "dashboards")
        return f.suggested, f.mentions

    assert feature(once + filler) == (False, 1)  # long document, named once: offered only
    assert feature(often + filler[2:])[0] is True  # named in three places: ticked
    assert feature(once)[0] is True  # short document: every mention counts


EXTRA_PAGES = [
    ["1. Scope", "1.1 This is a pilot project for the department's grievance portal."],
    [
        "2. Terms",
        "2.1 The vendor shall deploy 2 resident engineers at the department during O&M.",
        "2.2 Liquidated damages of 0.5% per week of delay shall apply.",
        "2.3 A warranty of one year applies after go-live.",
        "2.4 Payment terms: 30% on signing of the agreement, 50% on go-live and 20% after "
        "three months of stable operation.",
    ],
]


def test_basic_reader_lists_what_it_does_not_price(
    client: TestClient, admin_headers: H, no_override: None
) -> None:
    """Penalties, warranty and payment terms are shown with their page; onsite staff and an
    MVP hint are read and priced."""
    doc = upload(client, admin_headers, "terms.pdf", make_pdf(EXTRA_PAGES)).json()
    ex = doc["extraction"]
    topics = {n["topic"]: n for n in ex["not_priced"]}
    assert {"penalties", "warranty", "payment_terms"} <= set(topics)
    assert topics["warranty"]["evidence"][0]["page"] == 2
    assert topics["penalties"]["label"].startswith("Penalties")
    assert ex["numbers"]["onsite_staff"]["value"] == 2
    assert ex["choices"]["version"]["value"] == "mvp"  # "pilot project"
    body = {"name": "Grievances", "document_id": doc["id"], "features": ["grievances"]}
    without = client.post("/quick-estimates", json=body, headers=admin_headers).json()
    staffed = client.post(
        "/quick-estimates", json={**body, "overrides": {"onsite_staff": 2}}, headers=admin_headers
    ).json()
    items = {i["key"]: i for i in staffed["cost_items"]}
    assert items["onsite"]["amount_paise"] > 0
    assert staffed["likely_paise"] > without["likely_paise"]


def test_ai_reader_payment_terms_are_priced(client: TestClient, admin_headers: H) -> None:
    quote = "30% on signing of the agreement, 50% on go-live and 20% after"

    def responder(system: str, user: str) -> object:
        assert "penalties:" in system and "payment_terms" in system
        return {
            "project_name": "Grievance portal",
            "summary": "A grievance portal.",
            "features": [{"key": "grievances", "evidence": [{"page": 1, "clause_ref": "1.1", "quote": "pilot project for the department's grievance portal"}]}],
            "other_features": [],
            "choices": [],
            "numbers": [],
            "unclear_points": [],
            "not_priced": [
                {"topic": "warranty", "note": "One year of warranty.", "evidence": [{"page": 2, "clause_ref": "2.3", "quote": "A warranty of one year applies after go-live"}]},
                {"topic": "payment_terms", "note": "Paid in three parts.", "evidence": [{"page": 2, "clause_ref": "2.4", "quote": quote}]},
                {"topic": "spaceship", "note": "", "evidence": [{"page": 2, "clause_ref": "", "quote": quote}]},
            ],
            "payment_terms": {
                "payments": [
                    {"percent": 30, "when": "on_order"},
                    {"percent": 50, "when": "go_live"},
                    {"percent": 20, "when": "after_go_live", "months_after": 3},
                ],
                "evidence": [{"page": 2, "clause_ref": "2.4", "quote": quote}],
            },
        }  # fmt: skip

    set_llm_client(FakeClient(responder))
    try:
        doc = upload(client, admin_headers, "terms.pdf", make_pdf(EXTRA_PAGES)).json()
    finally:
        set_llm_client(None)
    ex = doc["extraction"]
    terms = ex["payment_terms"]
    assert [p["label"] for p in terms["payments"]] == [
        "30% on order",
        "50% go live",
        "20% after go live + 3 months",
    ]
    assert {n["topic"] for n in ex["not_priced"]} == {"warranty"}  # payment terms are priced
    schedule = [
        {k: p[k] for k in ("at_build_fraction", "offset_months", "percent")}
        for p in terms["payments"]
    ]
    body = {
        "name": "Grievances",
        "document_id": doc["id"],
        "features": ["grievances"],
        "payment_schedule": schedule,
    }
    r = client.post("/quick-estimates", json=body, headers=admin_headers)
    assert r.status_code == 201, r.text
    bad = [{**schedule[0], "percent": "0.9"}, *schedule[1:]]
    r = client.post(
        "/quick-estimates", json={**body, "payment_schedule": bad}, headers=admin_headers
    )
    assert r.status_code == 422  # shares must add up to 100%
