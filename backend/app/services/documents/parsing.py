"""Turn an uploaded tender document into page-numbered text (brief Section 9, step 1).

PDF: pdfplumber, one entry per physical page. DOCX: Word files have no fixed pages, so text is
split into approximate pages of `DOCX_CHARS_PER_PAGE` characters (marked `approximate`).
Scanned PDFs (no text layer) are marked `scanned`; the AI reader reads them as PDFs.
"""

import io
import re
from dataclasses import dataclass

DOCX_CHARS_PER_PAGE = 3000  # roughly one printed A4 page of tender text
MIN_TEXT_CHARS_PER_PAGE = 20  # below this a PDF page is treated as having no text layer
PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"


SCANNED_NEEDS_AI = (
    "This PDF looks scanned (no selectable text). Scanned documents can be read once an AI key "
    "is set up; until then please upload a text PDF or the Word version."
)


class DocumentError(ValueError):
    """The file cannot be read; the message is safe to show to the user."""


@dataclass(frozen=True)
class Page:
    number: int
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    kind: str  # "pdf" | "docx"
    pages: list[Page]
    approximate_pages: bool
    scanned: bool = False  # no text layer: only an AI that reads PDFs can extract it

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    name = filename.lower()
    if name.endswith(".pdf") and data.startswith(PDF_MAGIC):
        return _parse_pdf(data)
    if name.endswith(".docx") and data.startswith(ZIP_MAGIC):
        return _parse_docx(data)
    raise DocumentError("Please upload a PDF or a Word (.docx) file.")


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _parse_pdf(data: bytes) -> ParsedDocument:
    import pdfplumber

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [
                Page(number=i, text=_clean(page.extract_text() or ""))
                for i, page in enumerate(pdf.pages, start=1)
            ]
    except Exception as exc:
        raise DocumentError("This PDF could not be opened. Is it password-protected?") from exc
    if not pages:
        raise DocumentError("This PDF has no pages.")
    with_text = [p for p in pages if len(p.text) >= MIN_TEXT_CHARS_PER_PAGE]
    if not with_text:
        return ParsedDocument(kind="pdf", pages=pages, approximate_pages=False, scanned=True)
    return ParsedDocument(kind="pdf", pages=pages, approximate_pages=False)


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("This Word file could not be opened.") from exc
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(dict.fromkeys(cells)))
    pages: list[Page] = []
    current: list[str] = []
    size = 0
    for block in blocks:
        current.append(block)
        size += len(block)
        if size >= DOCX_CHARS_PER_PAGE:
            pages.append(Page(number=len(pages) + 1, text=_clean("\n".join(current))))
            current, size = [], 0
    if current:
        pages.append(Page(number=len(pages) + 1, text=_clean("\n".join(current))))
    if not pages:
        raise DocumentError("This Word file has no text.")
    return ParsedDocument(kind="docx", pages=pages, approximate_pages=True)


def normalise(text: str) -> str:
    """For citation checks: case-insensitive, whitespace- and quote-insensitive."""
    text = text.lower().replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    return re.sub(r"\s+", " ", re.sub(r"[-–—]", "-", text)).strip()
